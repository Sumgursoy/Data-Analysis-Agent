"""Faz 2 uçtan uca testi — kasten zorlaştırılmış gerçekçi dosyalar."""

from pathlib import Path

import pandas as pd
from fastapi.testclient import TestClient
from openpyxl import Workbook

from backend.main import app

TMP = Path(__file__).parent / "veri"
TMP.mkdir(exist_ok=True)


def kotu_excel() -> Path:
    """Kurumsal Excel: başlık 4. satırda, üstte logo/başlık, altta toplam,
    Türkçe sayı formatı, ikinci sheet, bir de boş sheet."""
    wb = Workbook()

    ws = wb.active
    ws.title = "Satışlar 2024"
    ws.append(["ACME A.Ş."])                     # 1: şirket başlığı
    ws.append([])                                # 2: boş
    ws.append(["Satış Raporu", None, "2024"])    # 3: rapor başlığı
    ws.append(["Müşteri No", "Sipariş Tarihi", "Bölge", "TUTAR(₺)", "Durum"])  # 4: BAŞLIK
    veri = [
        ("1001", "31.12.2024", "Ege", "1.234,56", "tamamlandı"),
        ("1002", "15.01.2024", "Marmara", "12.000,00", "tamamlandı"),
        ("1003", "03.02.2024", "Ege", "(500,00)", "iptal"),
        ("1004", "22.07.2024", "İç Anadolu", "8.750,25", "tamamlandı"),
        ("1005", "09.11.2024", "Marmara", "3.100,00", "beklemede"),
    ]
    for satir in veri:
        ws.append(list(satir))
    ws.append(["GENEL TOPLAM", None, None, "24.584,81", None])   # alt toplam

    ws2 = wb.create_sheet("Müşteriler")
    ws2.append(["musteri_no", "ad", "segment"])
    for no, ad, seg in [
        ("1001", "Yılmaz Ltd", "kurumsal"), ("1002", "Öz Ticaret", "kobi"),
        ("1003", "Şahin A.Ş.", "kurumsal"), ("1004", "Çelik San.", "kobi"),
        ("1005", "Güneş Ltd", "kobi"),
    ]:
        ws2.append([no, ad, seg])

    wb.create_sheet("Boş Sayfa")  # tamamen boş → elenmeli

    yol = TMP / "satislar_2024.xlsx"
    wb.save(yol)
    return yol


def cp1254_csv() -> Path:
    yol = TMP / "sube_ciro.csv"
    icerik = (
        "Şube Adı;Şehir;Ciro (TL);Açılış Tarihi\n"
        "Kadıköy Şubesi;İstanbul;1.250.000,50;01.03.2019\n"
        "Çankaya Şubesi;Ankara;980.400,00;15.06.2020\n"
        "Konak Şubesi;İzmir;1.100.750,25;22.09.2018\n"
        "Nilüfer Şubesi;Bursa;760.000,00;05.01.2021\n"
    )
    yol.write_bytes(icerik.encode("cp1254"))  # UTF-8 DEĞİL
    return yol


def tanimsiz_dosya() -> Path:
    yol = TMP / "kayit.dat"
    satirlar = ["#LOG v2.1 baslik"] + [
        f"2024-03-{g:02d}|OK|{g * 137}|islem-{g}" for g in range(1, 26)
    ]
    yol.write_text("\n".join(satirlar), encoding="utf-8")
    return yol


def main():
    with TestClient(app) as c:
        sid = c.post("/api/sessions").json()["session_id"]
        print(f"session: {sid}\n")

        for yol in (kotu_excel(), cp1254_csv(), tanimsiz_dosya()):
            with yol.open("rb") as f:
                r = c.post(
                    f"/api/sessions/{sid}/sources/file",
                    files={"file": (yol.name, f, "application/octet-stream")},
                )
            print(f"{'=' * 74}\n{yol.name}  →  HTTP {r.status_code}")
            if r.status_code != 200:
                print("  HATA:", r.text[:400])
                continue
            for ds in r.json()["datasets"]:
                print(f"\n--- {ds['name']} ({ds['kind']}) ---")
                print(ds["schema_card"])

        print(f"\n{'=' * 74}\nKATALOG (agent'in gordugu prompt blogu):\n")
        print(c.get(f"/api/sessions/{sid}/catalog").json()["prompt_block"])

        print("\n--- diskteki parquet dosyalari ---")
        from backend import config
        for p in sorted((config.session_dir(sid) / "data").iterdir()):
            print(f"  {p.name:<34} {p.stat().st_size:>9,} bayt")

        print("\n--- parquet gercekten dogru mu (tipler + toplam) ---")
        df = pd.read_parquet(config.session_dir(sid) / "data" / "satislar_2024.parquet")
        print(df.dtypes.to_string())
        print("tutar toplami:", df["tutar_tl"].sum(), " (beklenen 24584.81)")
        print("musteri_no ornek:", df["musteri_no"].tolist())
        print("tarih min/max:", df["siparis_tarihi"].min(), df["siparis_tarihi"].max())

        c.delete(f"/api/sessions/{sid}")


main()
